from __future__ import annotations

import hashlib
import json
import os
from io import BytesIO
from pathlib import Path
from typing import Any

import chromadb
import requests
from docx import Document
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from pypdf import PdfReader
from sentence_transformers import SentenceTransformer

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "chroma_db"
OLLAMA_URL = os.getenv("OLLAMA_URL", "http://localhost:11434")
LLM_MODEL = os.getenv("LLM_MODEL", "llama3.2:1b")
EMBED_MODEL = "all-MiniLM-L6-v2"
TOP_K = 3
MAX_DISTANCE = 0.65
ALLOWED_EXTENSIONS = {".pdf", ".txt", ".docx"}
MAX_UPLOAD_BYTES = 20 * 1024 * 1024  # 20 MB per file

app = FastAPI(title="InsightFlow API", version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Load the embedding model once when the API process starts.
model = SentenceTransformer(EMBED_MODEL)
client = chromadb.PersistentClient(path=str(DB_PATH))
collection = client.get_or_create_collection(
    name="my_documents", metadata={"hnsw:space": "cosine"}
)


class ChatRequest(BaseModel):
    question: str


def read_document(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    try:
        if ext == ".txt":
            return data.decode("utf-8", errors="replace")
        if ext == ".pdf":
            reader = PdfReader(BytesIO(data))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        if ext == ".docx":
            doc = Document(BytesIO(data))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            # Include table text too, as many reports store content in tables.
            for table in doc.tables:
                for row in table.rows:
                    paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
            return "\n".join(paragraphs)
    except Exception as exc:
        raise ValueError(f"Could not read {filename}: {exc}") from exc
    raise ValueError("Supported formats are PDF, TXT, and DOCX.")


def split_into_chunks(text: str, size: int = 100, overlap: int = 20) -> list[str]:
    words = text.split()
    if not words:
        return []
    step = max(1, size - overlap)
    return [
        " ".join(words[start:start + size])
        for start in range(0, len(words), step)
        if " ".join(words[start:start + size]).strip()
    ]


def get_processed_files() -> list[dict[str, Any]]:
    result: dict[str, int] = {}
    records = collection.get(include=["metadatas"])
    for meta in records.get("metadatas") or []:
        if meta:
            name = meta.get("source", "Unknown")
            result[name] = result.get(name, 0) + 1
    return [
        {"name": name, "chunks": chunks, "extension": Path(name).suffix.lower().lstrip(".")}
        for name, chunks in sorted(result.items(), key=lambda item: item[0].lower())
    ]


@app.get("/health")
def health():
    ollama_online = False
    try:
        ollama_online = requests.get(f"{OLLAMA_URL}/api/tags", timeout=2).ok
    except requests.RequestException:
        pass
    return {"status": "ok", "ollama_online": ollama_online, "indexed_chunks": collection.count()}


@app.get("/files")
def list_files():
    return {"files": get_processed_files()}


@app.post("/upload")
async def upload_files(files: list[UploadFile] = File(...)):
    results = []
    for upload in files:
        filename = Path(upload.filename or "document.txt").name
        ext = Path(filename).suffix.lower()
        if ext not in ALLOWED_EXTENSIONS:
            results.append({"name": filename, "status": "error", "message": "Only PDF, TXT, and DOCX are supported."})
            continue
        data = await upload.read()
        if len(data) > MAX_UPLOAD_BYTES:
            results.append({"name": filename, "status": "error", "message": "File is larger than 20 MB."})
            continue
        try:
            text = read_document(filename, data)
            if not text.strip():
                raise ValueError("No readable text found. Scanned PDFs need OCR and are not supported yet.")
            chunks = split_into_chunks(text)
            if not chunks:
                raise ValueError("No text chunks were created.")
            # Match the Streamlit app's behaviour: replace previous indexed version with the same filename.
            collection.delete(where={"source": filename})
            key = hashlib.sha256(filename.encode("utf-8")).hexdigest()[:16]
            embeddings = model.encode(chunks, normalize_embeddings=True, convert_to_numpy=True)
            collection.upsert(
                ids=[f"{key}_{i}" for i in range(len(chunks))],
                documents=chunks,
                embeddings=embeddings.tolist(),
                metadatas=[{"source": filename, "chunk_number": i + 1} for i in range(len(chunks))],
            )
            results.append({"name": filename, "status": "success", "chunks": len(chunks)})
        except Exception as exc:
            results.append({"name": filename, "status": "error", "message": str(exc)})
    return {"results": results, "files": get_processed_files()}


@app.delete("/files/{filename}")
def delete_file(filename: str):
    safe_name = Path(filename).name
    collection.delete(where={"source": safe_name})
    return {"ok": True, "files": get_processed_files()}


def retrieve(question: str) -> list[dict[str, Any]]:
    count = collection.count()
    if count == 0:
        return []
    embedding = model.encode(question, normalize_embeddings=True).tolist()
    result = collection.query(
        query_embeddings=[embedding],
        n_results=min(TOP_K, count),
        include=["documents", "distances", "metadatas"],
    )
    hits = []
    docs = (result.get("documents") or [[]])[0]
    distances = (result.get("distances") or [[]])[0]
    metadatas = (result.get("metadatas") or [[]])[0]
    for doc, distance, metadata in zip(docs, distances, metadatas):
        if distance <= MAX_DISTANCE:
            metadata = metadata or {}
            hits.append({
                "text": doc,
                "source": metadata.get("source", "Unknown"),
                "chunk": metadata.get("chunk_number", "?"),
                "match": max(0, min(100, round((1 - distance) * 100))),
            })
    return hits


def generate_answer(question: str, hits: list[dict[str, Any]]) -> str:
    if not hits:
        return "I couldn't find enough relevant information in your uploaded documents."
    context = "\n\n".join(
        f"[Source: {hit['source']}, part {hit['chunk']}]\n{hit['text']}" for hit in hits
    )
    prompt = f"""Answer the user's question using only the context below.
Be clear, useful, and concise. Do not invent facts or repeat points.
If the context does not contain the answer, say that the uploaded documents do not provide enough information.

Context:
{context}

Question: {question}

Answer:"""
    try:
        response = requests.post(
            f"{OLLAMA_URL}/api/generate",
            json={
                "model": LLM_MODEL,
                "prompt": prompt,
                "stream": False,
                "options": {"temperature": 0, "num_predict": 500},
            },
            timeout=600,
        )
        response.raise_for_status()
        answer = response.json().get("response", "").strip()
        return answer or "The model returned an empty answer. Please try again."
    except requests.RequestException as exc:
        raise HTTPException(
            status_code=503,
            detail=f"Could not contact Ollama. Make sure Ollama is running and model '{LLM_MODEL}' is installed. Details: {exc}",
        ) from exc


@app.post("/chat")
def chat(request: ChatRequest):
    question = request.question.strip()
    if not question:
        raise HTTPException(status_code=400, detail="Please enter a question.")
    if collection.count() == 0:
        raise HTTPException(status_code=400, detail="Upload a document first.")
    hits = retrieve(question)
    answer = generate_answer(question, hits)
    # Do not send full document passages unless the user opens source details in the UI.
    sources = [
        {
            "source": hit["source"],
            "chunk": hit["chunk"],
            "match": hit["match"],
            "snippet": hit["text"][:280] + ("…" if len(hit["text"]) > 280 else ""),
        }
        for hit in hits
    ]
    return {"answer": answer, "sources": sources}
